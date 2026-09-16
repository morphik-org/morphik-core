r"""Build a searchable evidence index and linked incident timeline.

Run from this repository or directly with:
    python build_incident_report.py "C:\Users\nadar\Downloads\EEOC 072026"
"""

from __future__ import annotations

import csv
import hashlib
import html
import json
import mimetypes
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path


SUPPORTED_TEXT = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml", ".log"}
SKIP_EXTENSIONS = {".db", ".sqlite", ".sqlite3", ".m4a", ".mp3", ".wav", ".png", ".jpg", ".jpeg", ".gif", ".zip"}
MAX_TEXT = 120_000
MAX_SNIPPET = 700


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in SKIP_EXTENSIONS:
        return ""
    try:
        if suffix in SUPPORTED_TEXT:
            return path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT]
        if suffix == ".pdf":
            from pypdf import PdfReader

            pages = []
            for page in PdfReader(str(path)).pages:
                pages.append(page.extract_text() or "")
                if sum(len(p) for p in pages) >= MAX_TEXT:
                    break
            return "\n\n".join(pages)[:MAX_TEXT]
        if suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)[:MAX_TEXT]
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            book = load_workbook(str(path), read_only=True, data_only=True)
            chunks = []
            for sheet in book.worksheets:
                chunks.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        chunks.append(" | ".join(values))
                    if sum(len(c) for c in chunks) >= MAX_TEXT:
                        break
                if sum(len(c) for c in chunks) >= MAX_TEXT:
                    break
            return "\n".join(chunks)[:MAX_TEXT]
    except Exception as exc:
        return f"[Extraction error: {type(exc).__name__}: {exc}]"
    return ""


def first_date(text: str, name: str, mtime: datetime) -> str:
    candidates = re.findall(r"\b20\d{2}[-_]\d{1,2}[-_]\d{1,2}\b", name + " " + text[:5000])
    normalized = []
    for value in candidates:
        try:
            normalized.append(datetime.strptime(value.replace("_", "-"), "%Y-%m-%d").date().isoformat())
        except ValueError:
            pass
    if normalized:
        return min(normalized)
    for value in re.findall(r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},?\s+20\d{2}\b", name + " " + text[:5000], re.I):
        for fmt in ("%B %d, %Y", "%B %d %Y"):
            try:
                return datetime.strptime(value.replace(",", ""), fmt).date().isoformat()
            except ValueError:
                continue
    return mtime.date().isoformat()


def doc_type(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".pdf": "PDF",
        ".docx": "Word document",
        ".xlsx": "Excel workbook",
        ".xlsm": "Excel workbook",
        ".csv": "CSV data",
        ".md": "Markdown",
        ".json": "JSON data",
        ".m4a": "Audio recording",
        ".png": "Image",
        ".jpg": "Image",
        ".jpeg": "Image",
    }.get(suffix, suffix[1:].upper() if suffix else "File")


def classify(path: Path, text: str) -> tuple[str, str, str, str]:
    blob = f"{path.name} {path.parent.name} {text[:10000]}".lower()
    if any(x in blob for x in ("eeoc", "charge", "agency", "online inquiry", "vec", "unemployment")):
        topic = "EEOC / agency / unemployment"
    elif any(x in blob for x in ("race", "racis", "racial", "black", "identity", "hostile work")):
        topic = "Race / hostile work environment"
    elif any(x in blob for x in ("accommodat", "ada", "disability", "medical", "doctor", "leave", "return to work")):
        topic = "Disability / accommodation / medical"
    elif any(x in blob for x in ("terminat", "retaliat", "resign", "separation", "position statement")):
        topic = "Retaliation / termination"
    elif any(x in blob for x in ("wage", "benefit", "damages", "unemployment", "salary", "pay")):
        topic = "Damages / benefits"
    elif any(x in blob for x in ("meeting", "call", "email", "message", "chat", "transcript", "communication")):
        topic = "Communications / meetings"
    elif any(x in blob for x in ("legal", "litigation", "discovery", "pleading", "evidence", "chronology")):
        topic = "Case preparation / legal"
    else:
        topic = "Other / uncategorized"

    if any(x in blob for x in ("email", "e-mail")):
        means = "Email"
    elif any(x in blob for x in ("call", "phone", "telephone", "voicemail")):
        means = "Phone call"
    elif any(x in blob for x in ("chat", "message", "sms", "text")):
        means = "Chat / text message"
    elif any(x in blob for x in ("meeting", "zoom", "transcript")):
        means = "Meeting"
    elif path.suffix.lower() in {".m4a", ".mp3", ".wav"}:
        means = "Audio recording"
    else:
        means = "Written record"

    names = [
        "Nada Boris", "Boris", "Goodwin Living", "Goodwin House", "Fran Casey",
        "Lindsay Hutter", "Kathie Miller", "Trish Povlitz", "EEOC", "VEC",
    ]
    participants = ", ".join(name for name in names if name.lower() in blob)
    if not participants:
        participants = "Not identified in extracted text"
    entities = []
    for entity in ("Goodwin Living", "Goodwin House", "EEOC", "Virginia Employment Commission", "VEC", "Dr. Gloria Okereke"):
        if entity.lower() in blob:
            entities.append(entity)
    org = ", ".join(dict.fromkeys(entities)) or "Not identified"
    return topic, means, participants, org


def safe_name(relative: Path, index: int) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", relative.name)
    return f"{index:04d}_{stem}.txt"


def link(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def build(root: Path) -> None:
    root = root.resolve()
    output_dir = root / "INCIDENT_REPORT_TEXT"
    duplicate_dir = root / "OLD_DUPLICATES"
    output_dir.mkdir(exist_ok=True)
    duplicate_dir.mkdir(exist_ok=True)

    excluded = {output_dir.resolve(), duplicate_dir.resolve()}
    files = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        relative_parts = path.relative_to(root).parts
        if "2026 09 16 updates" in relative_parts:
            continue
        if any(parent.resolve() in excluded for parent in [path, *path.parents]):
            continue
        if path.name in {"INCIDENT_REPORT_INDEX.html", "INCIDENT_TIMELINE.html", "INCIDENT_REPORT_DATA.json"}:
            continue
        files.append(path)

    groups: dict[str, list[Path]] = defaultdict(list)
    for path in files:
        try:
            groups[sha256(path)].append(path)
        except OSError:
            continue

    duplicate_moves = []
    canonical = set(files)
    for digest, matches in groups.items():
        if len(matches) < 2:
            continue
        keep = max(matches, key=lambda p: p.stat().st_mtime)
        for old in matches:
            if old == keep:
                continue
            destination = duplicate_dir / old.relative_to(root)
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists():
                destination = duplicate_dir / f"{digest[:12]}_{old.name}"
            shutil.move(str(old), str(destination))
            canonical.discard(old)
            duplicate_moves.append({"from": str(old.relative_to(root)), "to": str(destination.relative_to(root)), "sha256": digest})

    records = []
    for index, path in enumerate(sorted(canonical, key=lambda p: str(p).lower()), 1):
        relative = path.relative_to(root)
        stat = path.stat()
        extracted = extract_text(path)
        topic, means, participants, org = classify(path, extracted)
        text_file = ""
        if extracted:
            target = output_dir / safe_name(relative, index)
            target.write_text(extracted, encoding="utf-8")
            text_file = link(target, root)
        record = {
            "id": f"DOC-{index:04d}",
            "name": path.name,
            "path": relative.as_posix(),
            "link": link(path, root),
            "date": first_date(extracted, path.name, datetime.fromtimestamp(stat.st_mtime)),
            "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            "type": doc_type(path),
            "size": stat.st_size,
            "topic": topic,
            "means": means,
            "participants": participants,
            "org": org,
            "text_file": text_file,
            "content": re.sub(r"\s+", " ", extracted).strip()[:MAX_SNIPPET] if extracted else "No text extracted (binary, audio, image, database, or unsupported format).",
        }
        records.append(record)

    data = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "root": str(root),
        "document_count": len(records),
        "duplicate_moves": duplicate_moves,
        "documents": records,
    }
    (root / "INCIDENT_REPORT_DATA.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    formal_events = []
    timeline_csv = root / "Evidence Index" / "05 - Timeline March to June 2025 (Detailed).csv"
    if timeline_csv.exists():
        with timeline_csv.open(encoding="utf-8-sig", newline="") as fh:
            for row in csv.DictReader(fh):
                if row.get("Date") and row.get("What Happened"):
                    formal_events.append(row)

    by_topic = defaultdict(list)
    for record in records:
        by_topic[record["topic"]].append(record)
    categories = sorted(set(by_topic) | {row.get("Why This Step Matters", "Formal timeline") for row in formal_events})

    def page_shell(title: str, body: str, script: str = "") -> str:
        return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title>
<style>
body{{font:14px system-ui,Segoe UI,sans-serif;margin:24px;color:#202124;background:#f7f8fa}}
h1{{margin-bottom:4px}} .muted{{color:#5f6368}} nav a{{margin-right:16px}} .controls{{position:sticky;top:0;background:#fff;padding:12px;border:1px solid #ddd;z-index:2}}
input,select{{padding:8px;margin:4px;border:1px solid #bbb;border-radius:4px}} table{{border-collapse:collapse;width:100%;background:#fff}}
th,td{{border:1px solid #d8dbe0;padding:7px;vertical-align:top;text-align:left}} th{{background:#263238;color:#fff;position:sticky;top:75px}}
td.content{{max-width:460px;white-space:normal}} .tag{{display:inline-block;background:#e8eef7;border-radius:12px;padding:2px 7px;margin:1px}}
details{{background:#fff;border:1px solid #ddd;padding:10px;margin:8px 0}} summary{{cursor:pointer;font-weight:600}}
</style></head><body>{body}{script}</body></html>"""

    index_rows = []
    for r in records:
        index_rows.append(
            "<tr data-search=\"{search}\" data-topic=\"{data_topic}\" data-date=\"{date}\">"
            "<td>{id}</td><td>{date}</td><td><a href=\"{link}\">{name}</a><br><span class=\"muted\">{path}</span></td>"
            "<td>{type}</td><td>{participants}</td><td>{topic}</td><td>{means}</td><td>{org}</td>"
            "<td class=\"content\">{content}<br>{text_link}</td></tr>".format(
                search=html.escape(" ".join(str(v) for v in r.values()), quote=True),
                data_topic=html.escape(r["topic"], quote=True), id=r["id"], date=r["date"],
                link=html.escape(r["link"], quote=True), name=html.escape(r["name"]),
                path=html.escape(r["path"]), type=html.escape(r["type"]),
                participants=html.escape(r["participants"]), topic=html.escape(r["topic"]),
                means=html.escape(r["means"]), org=html.escape(r["org"]),
                content=html.escape(r["content"]),
                text_link=f'<a href="{html.escape(r["text_file"], quote=True)}">full extracted text</a>' if r["text_file"] else "",
            )
        )
    topics = "".join(f"<option>{html.escape(t)}</option>" for t in sorted(by_topic))
    index_body = f"""<h1>EEOC Evidence Document Index</h1>
<p class="muted">Generated {html.escape(data["generated"])} | {len(records)} indexed files | {len(duplicate_moves)} exact duplicates moved to <a href="OLD_DUPLICATES/">OLD_DUPLICATES</a>.</p>
<nav><a href="INCIDENT_TIMELINE.html">Open incident timeline</a><a href="INCIDENT_REPORT_DATA.json">Download index data</a></nav>
<div class="controls"><label>Search <input id="q" size="45" placeholder="name, participant, topic, content..."></label>
<label>Topic <select id="topic"><option value="">All topics</option>{topics}</select></label>
<label>Date order <select id="dateOrder"><option value="desc">Newest to oldest</option><option value="asc">Oldest to newest</option></select></label><span id="count"></span></div>
<table id="docs"><thead><tr><th>ID</th><th>Date</th><th>Document</th><th>Type</th><th>Participants</th><th>Topic</th><th>Means</th><th>Org/entity</th><th>Content / retrieval</th></tr></thead>
<tbody>{"".join(index_rows)}</tbody></table>"""
    index_script = """<script>
const body=document.querySelector('#docs tbody'),rows=[...body.querySelectorAll('tr')],q=document.querySelector('#q'),topic=document.querySelector('#topic'),dateOrder=document.querySelector('#dateOrder'),count=document.querySelector('#count');
function sortRows(){rows.sort((a,b)=>dateOrder.value==='desc'?b.dataset.date.localeCompare(a.dataset.date):a.dataset.date.localeCompare(b.dataset.date));rows.forEach(r=>body.appendChild(r));}
function apply(){sortRows();const needle=q.value.toLowerCase(), selected=topic.value;let n=0;rows.forEach(r=>{const ok=(!needle||r.dataset.search.toLowerCase().includes(needle))&&(!selected||r.dataset.topic===selected);r.hidden=!ok;if(ok)n++});count.textContent=` ${n} matching documents`;} q.oninput=apply;topic.onchange=apply;dateOrder.onchange=apply;apply();
</script>"""
    (root / "INCIDENT_REPORT_INDEX.html").write_text(page_shell("EEOC Evidence Document Index", index_body, index_script), encoding="utf-8")

    timeline_items = []
    for row in formal_events:
        source_text = row.get("Source (Document + Page)", "")
        sources = [r for r in records if any(token.lower() in (r["name"] + " " + r["path"]).lower() for token in re.findall(r"DOC-\d+", source_text))]
        links = " ".join(f'<a href="{html.escape(s["link"], quote=True)}">{html.escape(s["id"])}</a>' for s in sources[:8])
        category = row.get("Why This Step Matters", "Formal timeline")
        timeline_items.append({"date": row["Date"], "category": category, "event": row["What Happened"], "source": source_text, "links": links})
    for topic, topic_records in by_topic.items():
        for r in topic_records:
            timeline_items.append({"date": r["date"], "category": topic, "event": r["content"], "source": r["path"], "links": f'<a href="{html.escape(r["link"], quote=True)}">{html.escape(r["id"])}</a>'})
    timeline_items.sort(key=lambda x: (x["date"], x["category"], x["event"]))
    timeline_rows = "".join(
        f'<tr data-search="{html.escape((i["category"]+" "+i["event"]+" "+i["source"]).lower(), quote=True)}" data-date="{html.escape(i["date"], quote=True)}">'
        f'<td>{html.escape(i["date"])}</td><td><span class="tag">{html.escape(i["category"])}</span></td>'
        f'<td>{html.escape(i["event"])}</td><td>{html.escape(i["source"])}<br>{i["links"]}</td></tr>'
        for i in timeline_items
    )
    category_links = "".join(f'<a href="#cat-{n}">{html.escape(n)}</a> ' for n in sorted(by_topic))
    category_details = []
    for topic in sorted(by_topic):
        links = " ".join(
            f'<a href="{html.escape(record["link"], quote=True)}">{html.escape(record["id"])} {html.escape(record["name"])}</a>'
            for record in by_topic[topic]
        )
        category_details.append(
            f'<details id="cat-{html.escape(topic)}"><summary>{html.escape(topic)} '
            f'({len(by_topic[topic])} documents)</summary><p>{links}</p></details>'
        )
    timeline_body = f"""<h1>Mass Incident Timeline</h1><p class="muted">Formal chronology plus document-linked incident records. Each row links to its supporting evidence.</p>
<nav><a href="INCIDENT_REPORT_INDEX.html">Back to document index</a> {category_links}</nav>
<div class="controls"><label>Search <input id="timelineQ" size="55" placeholder="incident, date, person, source..."></label>
<label>Date order <select id="timelineDateOrder"><option value="desc">Newest to oldest</option><option value="asc">Oldest to newest</option></select></label><span id="timelineCount"></span></div>
<table id="timeline"><thead><tr><th>Date</th><th>Incident category</th><th>Event / document content</th><th>Supporting document(s)</th></tr></thead><tbody>{timeline_rows}</tbody></table>
<h2>Incident categories</h2>{"".join(category_details)}"""
    timeline_script = """<script>
const timelineBody=document.querySelector('#timeline tbody'),trs=[...timelineBody.querySelectorAll('tr')],tq=document.querySelector('#timelineQ'),dateOrder=document.querySelector('#timelineDateOrder'),tc=document.querySelector('#timelineCount');
function sortTimeline(){trs.sort((a,b)=>dateOrder.value==='desc'?b.dataset.date.localeCompare(a.dataset.date):a.dataset.date.localeCompare(b.dataset.date));trs.forEach(r=>timelineBody.appendChild(r));}
function filterTimeline(){sortTimeline();const n=tq.value.toLowerCase();let c=0;trs.forEach(r=>{const ok=!n||r.dataset.search.includes(n);r.hidden=!ok;if(ok)c++});tc.textContent=` ${c} matching incidents`;} tq.oninput=filterTimeline;dateOrder.onchange=filterTimeline;filterTimeline();
</script>"""
    (root / "INCIDENT_TIMELINE.html").write_text(page_shell("Mass Incident Timeline", timeline_body, timeline_script), encoding="utf-8")
    print(json.dumps({"documents": len(records), "timeline_events": len(timeline_items), "duplicates_moved": len(duplicate_moves), "index": str(root / "INCIDENT_REPORT_INDEX.html"), "timeline": str(root / "INCIDENT_TIMELINE.html")}, indent=2))


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\nadar\Downloads\EEOC 072026"))
